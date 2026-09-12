"""Build deck/smart-watchdog-deck.pptx natively (python-pptx): fixed title bar, centered body, uncropped screenshots, notes."""
import json, sqlite3, pathlib, re
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image

ROOT = pathlib.Path(__file__).resolve().parent.parent
IMG = ROOT / "docs" / "img"
ACC, MUTED, BORDER, TEXT = RGBColor(0x2d,0x4a,0x3e), RGBColor(0x6b,0x72,0x80), RGBColor(0xe5,0xe7,0xeb), RGBColor(0x0a,0x0a,0x0a)
FONT = "PingFang TC"
W, H = Inches(13.333), Inches(7.5)
prs = Presentation(); prs.slide_width, prs.slide_height = W, H
blank = prs.slide_layouts[6]
notes = re.split(r"^## \d+\n", (ROOT/"deck"/"speaker-notes.md").read_text(), flags=re.M)[1:]

con = sqlite3.connect(ROOT/"data"/"watchdog.sqlite")
bt = con.execute("SELECT obs_year,n_obs,n_pos,ROUND(auc,3),metrics FROM app_backtests WHERE model_id=9 AND baseline IS NULL ORDER BY obs_year").fetchall()
cnt = {y:c for y,c in con.execute("SELECT obs_year, ROUND(auc,3) FROM app_backtests WHERE model_id=9 AND baseline='count'")}
lv = dict(con.execute("SELECT level, COUNT(*) FROM v_scores GROUP BY level").fetchall())
mid = {"tp":0,"fp":0,"fn":0,"tn":0}
for r in bt:
    m = json.loads(r[4])["at"]["mid"]
    for k in mid: mid[k] += m[k]
rec = mid["tp"]/(mid["tp"]+mid["fn"]); prec = mid["tp"]/(mid["tp"]+mid["fp"]); acc = (mid["tp"]+mid["tn"])/sum(mid.values())

def tb(slide, x, y, w, h, text, size=18, bold=False, color=TEXT, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font=FONT):
    box = slide.shapes.add_textbox(x, y, w, h); tf = box.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05); tf.margin_top = tf.margin_bottom = Inches(0.03)
    lines = text if isinstance(text, list) else [text]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.alignment = align
        run = p.add_run(); run.text = line; run.font.size = Pt(size); run.font.bold = bold; run.font.color.rgb = color; run.font.name = font
    return box

def bullets(slide, x, y, w, h, items, size=16, anchor=MSO_ANCHOR.MIDDLE):
    box = slide.shapes.add_textbox(x, y, w, h); tf = box.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.space_after = Pt(8)
        parts = re.split(r"(\*\*.*?\*\*)", it)
        r0 = p.add_run(); r0.text = "•  "; r0.font.size = Pt(size); r0.font.color.rgb = ACC; r0.font.name = FONT
        for part in parts:
            if not part: continue
            r = p.add_run(); r.text = part.strip("*"); r.font.size = Pt(size); r.font.bold = part.startswith("**"); r.font.color.rgb = TEXT; r.font.name = FONT
    return box

def title(slide, text):
    tb(slide, Inches(0.6), Inches(0.35), W - Inches(1.2), Inches(0.8), text, size=30, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

def foot(slide, n, dark=False):
    c = RGBColor(0xdd,0xe5,0xdf) if dark else MUTED
    tb(slide, Inches(0.6), H - Inches(0.45), Inches(9), Inches(0.3), "小小守護員 Smart Watchdog · 新北市教育局 AI 黑客松 · 資料日期 2026-09-11 · 事件層數字", size=10, color=c)
    tb(slide, W - Inches(1.6), H - Inches(0.45), Inches(1), Inches(0.3), str(n), size=10, color=c, align=PP_ALIGN.RIGHT)

def picture_fit(slide, name, x, y, w, h):
    im = Image.open(IMG/name); iw, ih = im.size; s = min(w/iw, h/ih); pw, ph = int(iw*s), int(ih*s)
    pic = slide.shapes.add_picture(str(IMG/name), x + (w-pw)//2, y + (h-ph)//2, pw, ph)
    pic.line.color.rgb = BORDER; pic.line.width = Pt(0.75)

def kpi(slide, x, y, w, h, v, l):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h); s.fill.solid(); s.fill.fore_color.rgb = RGBColor(0xff,0xff,0xff); s.line.color.rgb = BORDER; s.shadow.inherit = False
    tb(slide, x + Inches(0.15), y + Inches(0.12), w - Inches(0.3), Inches(0.6), v, size=28, bold=True, color=ACC, font="Menlo")
    tb(slide, x + Inches(0.15), y + Inches(0.75), w - Inches(0.3), h - Inches(0.8), l, size=13, color=MUTED)

BODY_Y, BODY_H = Inches(1.3), H - Inches(1.3) - Inches(0.6)
def shot_slide(n, ttl, img, items, note):
    s = prs.slides.add_slide(blank); title(s, ttl)
    picture_fit(s, img, Inches(0.6), BODY_Y, Inches(8.1), BODY_H)
    bullets(s, Inches(8.9), BODY_Y, Inches(3.9), BODY_H, items); foot(s, n); s.notes_slide.notes_text_frame.text = note

# 1 cover
s = prs.slides.add_slide(blank); bg = s.background.fill; bg.solid(); bg.fore_color.rgb = ACC
tb(s, Inches(0.8), Inches(1.6), Inches(11), Inches(0.5), "新北市政府 AI 黑客松 · 教育局命題", size=16, color=RGBColor(0xdd,0xe5,0xdf))
tb(s, Inches(0.8), Inches(2.1), Inches(11), Inches(1.3), "小小守護員", size=64, bold=True, color=RGBColor(0xff,0xff,0xff))
tb(s, Inches(0.8), Inches(3.35), Inches(11), Inches(0.9), "誰有可能是高風險幼兒園", size=36, color=RGBColor(0xff,0xff,0xff))
tb(s, Inches(0.8), Inches(4.45), Inches(11), Inches(0.6), "AI × 鑑識會計 · 教保機構智慧風險預警管理系統", size=22, color=RGBColor(0xff,0xff,0xff))
tb(s, Inches(0.8), Inches(5.1), Inches(11), Inches(0.6), "一份名單，三個答案：這季該去哪、為什麼、跑得完", size=18, color=RGBColor(0xdd,0xe5,0xdf))
foot(s, 1, dark=True); s.notes_slide.notes_text_frame.text = notes[0]
# 2
s = prs.slides.add_slide(blank); title(s, "本季高風險名單：都是回頭客")
picture_fit(s, "season.png", Inches(0.6), BODY_Y, Inches(8.1), BODY_H)
kpi(s, Inches(8.9), Inches(1.6), Inches(1.85), Inches(1.5), str(lv.get("高",0)), "高風險（再犯機率 ≥ 22%）"); kpi(s, Inches(10.95), Inches(1.6), Inches(1.85), Inches(1.5), str(lv.get("中",0)), "中風險（≥ 10%）")
bullets(s, Inches(8.9), Inches(3.3), Inches(3.9), Inches(3.2), ["近 12 個月有裁罰的 74 家必訪，名下同負責人的 15 家待人工確認","每一列寫著：再犯機率、近 36 月違規類型、到園要看的三件事","這不是猜的，是回頭客資料算出來的"], anchor=MSO_ANCHOR.TOP)
foot(s, 2); s.notes_slide.notes_text_frame.text = notes[1]
# 3
s = prs.slides.add_slide(blank); title(s, "問題在回頭客，不在大海撈針")
for i,(v,l) in enumerate([("48%","有裁罰園再犯（234 / 487）"),("75%","事件來自回頭客"),("63%","兩次裁罰間隔在一年內"),("26–33%","裁罰後 12 個月內再犯")]):
    kpi(s, Inches(0.6 + i*3.1), Inches(2.0), Inches(2.9), Inches(1.5), v, l)
tb(s, Inches(0.8), Inches(3.9), Inches(11.7), Inches(1.2), "一家幼兒園被罰完的那一天，承辦通常鬆一口氣結案；資料說，接下來 12 個月正是它最可能再出事的時候。", size=24, bold=True, color=ACC, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
tb(s, Inches(0.8), Inches(5.2), Inches(11.7), Inches(0.4), "新北裁罰 1,474 列 → 去重為 1,004 個園×日期事件；全國教保資訊網公開資料", size=12, color=MUTED, align=PP_ALIGN.CENTER)
foot(s, 3); s.notes_slide.notes_text_frame.text = notes[2]
# 4
shot_slide(4, "單一園所怎麼看：結論在上，證據在下", "detail-top.png", ["**第一眼**：分數、12 個月再犯機率、全市排名、裁罰事件數","**分數怎麼來**：模型對每個特徵的貢獻，正的推高、負的拉低","**時間軸**寫具體違反什麼——不當對待、超收、師生比——不寫條號","往下是負責人關聯圖、財務體檢、輿情"], notes[3])
# 5 method
s = prs.slides.add_slide(blank); title(s, "分數怎麼來：AUC 選模型，recall 定門檻")
tb(s, Inches(0.6), Inches(1.45), Inches(4.2), Inches(0.4), "五年 walk-forward", size=18, bold=True, color=ACC)
rows = len(bt)+1; tbl = s.shapes.add_table(rows, 5, Inches(0.6), Inches(1.9), Inches(4.2), Inches(0.36)*rows).table
hdr = ["測試年","觀察點","再犯","模型 AUC","按次數"]
for j,hh in enumerate(hdr):
    c = tbl.cell(0,j); c.text = hh; p = c.text_frame.paragraphs[0]; p.runs[0].font.size = Pt(11); p.runs[0].font.color.rgb = MUTED; p.runs[0].font.name = FONT; c.fill.solid(); c.fill.fore_color.rgb = RGBColor(0xf3,0xf4,0xf6)
for i,r in enumerate(bt, start=1):
    vals = [str(r[0]), f"{r[1]:,}", str(r[2]), f"{r[3]:.3f}", f"{cnt.get(r[0],0):.3f}"]
    for j,v in enumerate(vals):
        c = tbl.cell(i,j); c.text = v; p = c.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT if j else PP_ALIGN.LEFT
        p.runs[0].font.size = Pt(12); p.runs[0].font.bold = (j == 3); p.runs[0].font.name = "Menlo" if j else FONT; c.fill.solid(); c.fill.fore_color.rgb = RGBColor(0xff,0xff,0xff)
tb(s, Inches(0.6), Inches(4.3), Inches(4.2), Inches(0.9), "測試年只用前一年之前的資料訓練（365 天空窗）；邏輯迴歸、13 個事件史特徵、seed 固定", size=11, color=MUTED)
tb(s, Inches(5.1), Inches(1.45), Inches(3.8), Inches(0.4), "抽檢寧可錯殺", size=18, bold=True, color=ACC)
cells = [("", None), ("預測再犯", None), ("預測不再犯", None), ("實際再犯", None), (f"{mid['tp']:,}\n抓到", RGBColor(0xe8,0xef,0xea)), (f"{mid['fn']:,}\n漏抓", RGBColor(0xfe,0xe2,0xe2)), ("實際未再犯", None), (f"{mid['fp']:,}\n白跑", RGBColor(0xfe,0xf3,0xc7)), (f"{mid['tn']:,}", RGBColor(0xf3,0xf4,0xf6))]
cw, ch = Inches(1.25), Inches(0.85)
for k,(t,fill) in enumerate(cells):
    r_, c_ = divmod(k, 3); x = Inches(5.1) + c_*cw; y = Inches(1.95) + r_*ch
    if fill:
        sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, cw - Inches(0.05), ch - Inches(0.05)); sh.fill.solid(); sh.fill.fore_color.rgb = fill; sh.line.fill.background(); sh.shadow.inherit = False
        tf = sh.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE; a, b = t.split("\n") if "\n" in t else (t, "")
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; r = p.add_run(); r.text = a; r.font.size = Pt(20); r.font.bold = True; r.font.color.rgb = TEXT; r.font.name = "Menlo"
        if b: p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER; r2 = p2.add_run(); r2.text = b; r2.font.size = Pt(10); r2.font.color.rgb = MUTED; r2.font.name = FONT
    elif t:
        tb(s, x, y, cw, ch, t, size=11, color=MUTED, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
tb(s, Inches(5.1), Inches(4.6), Inches(3.8), Inches(0.4), "門檻「中」＝再犯機率 ≥ 10%，五年匯總", size=11, color=MUTED)
tb(s, Inches(9.2), Inches(1.45), Inches(3.6), Inches(0.4), "怎麼讀", size=18, bold=True, color=ACC)
bullets(s, Inches(9.2), Inches(1.9), Inches(3.6), Inches(4.2), [f"**Recall {rec:.2f}**：會再犯的園，七成被抓進名單", f"Precision {prec:.2f}：五家白跑四家，那是抽檢的代價", f"Accuracy {acc:.2f} 沒有意義：基準率 15%，什麼都不抓 accuracy 就 85%", "門檻設在人力跑得完的最低點；人力變了，設定頁改"], size=14, anchor=MSO_ANCHOR.TOP)
foot(s, 5); s.notes_slide.notes_text_frame.text = notes[4]
# 6 map
s = prs.slides.add_slide(blank); title(s, "全市在哪：每園一個點，點色＝再犯風險"); picture_fit(s, "overview.png", Inches(0.6), BODY_Y, W - Inches(1.2), BODY_H); foot(s, 6); s.notes_slide.notes_text_frame.text = notes[5]
# 7–9, 11
shot_slide(7, "跑得完：人力是約束，目標可以選", "schedule.png", ["3 位稽查員 × 每週 8 次 × 13 週 = 312 次產能，CP-SAT 排出週×稽查員行程","名單必訪、各區配額、釘選／排除都是硬約束；不可行時直接說原因","三種目標：風險優先／同區同負責人併訪／各區均衡","**歷史重播**：同樣 312 次，命中次年被罰園 40–53%，輪流只有 25–28%"], notes[6])
shot_slide(8, "同負責人：名下其他園列入待確認", "linker.png", ["負責人是公開登記資料；名下多園時，任一園被罰，其他園進「待確認」層","只列名單、**不進分數**：資料顯示持多園本身不是風險因子","同名不同人可一鍵排除，排名、匯出、關聯圖同步"], notes[7])
shot_slide(9, "鑑識會計是附錄，不是預言", "detail-finance.png", ["38 家非營利園、155 園年決算，OCR 後過資產負債恆等式","三燈：人事費率、每核定名額收入、內控查核表；門檻用全體分佈百分位定義","被罰與未罰園的比率幾乎一樣（人事費率 0.58 vs 0.56）→ **不進分數**","它回答的是「補助的錢有沒有花對」"], notes[8])
# 10
s = prs.slides.add_slide(blank); title(s, "不是又一個儀表板")
for i,(h,t) in enumerate([("根因在名單裡","每列附違規類型計數與稽查重點三行，稽查員到園前就知道要看什麼"),("人力算得出","產能是系統內的約束，行程用最佳化排出，歷史重播給命中率"),("限制講在前面","無紀錄不等於安全、財報不預測裁罰、precision 0.17 是抽檢的代價——寫在畫面上")]):
    x = Inches(0.6 + i*4.15); tb(s, x, Inches(2.4), Inches(3.9), Inches(0.8), h, size=28, bold=True, color=ACC); tb(s, x, Inches(3.3), Inches(3.9), Inches(2), t, size=16)
foot(s, 10); s.notes_slide.notes_text_frame.text = notes[9]
shot_slide(11, "全部公開、全部可重跑", "quality.png", ["資料來源：全國教保資訊網、新北幼教資源網決算 PDF、Google 新聞","裁罰去重到事件層；被罰個人姓名在資料層就擋掉","**make bootstrap**：clone 到真資料畫面五分鐘，不需 PDF、OCR、API key","35 條自動測試；模型核准與撤銷都留痕"], notes[10])
# 12
s = prs.slides.add_slide(blank); title(s, "限制，與命題三指標")
for i,(v,l) in enumerate([("風險辨識率","中門檻 recall 0.72；模型 AUC 0.64 vs 按次數 0.61"),("縮短預警時間","觀察點從裁罰後起算；下一次事件中位 151 天後，季排程趕得上"),("降低人力負擔","同樣 312 次／季，命中 40–53% vs 輪流 25–28%")]):
    x = Inches(0.6 + i*4.15); sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, Inches(1.6), Inches(3.9), Inches(1.7)); sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor(0xff,0xff,0xff); sh.line.color.rgb = BORDER; sh.shadow.inherit = False
    tb(s, x + Inches(0.15), Inches(1.7), Inches(3.6), Inches(0.6), v, size=22, bold=True, color=ACC); tb(s, x + Inches(0.15), Inches(2.3), Inches(3.6), Inches(1), l, size=13, color=MUTED)
bullets(s, Inches(0.6), Inches(3.6), Inches(12), Inches(2.2), ["預測的是「12 個月內再被裁罰」，不是傷害；710 家「無紀錄」不等於安全","財報只有公共化園；輿情會混入同名園，只做情蒐","precision 0.17：名單是優先順序，不是有罪名單"], size=16, anchor=MSO_ANCHOR.TOP)
tb(s, Inches(0.6), Inches(6.1), Inches(12), Inches(0.4), "GitHub mumigood/NTPC-childcare-risk-monitor · docs/user-manual.md", size=12, color=MUTED)
foot(s, 12); s.notes_slide.notes_text_frame.text = notes[11]
out = ROOT/"deck"/"smart-watchdog-deck.pptx"; prs.save(out); print(out, round(out.stat().st_size/1e6, 1), "MB", len(prs.slides), "slides")
